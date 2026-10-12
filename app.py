"""
Portfolio chatbot: RAG (Retrieval-Augmented Generation) with LangChain.

Read top to bottom, it runs in 5 steps:
  1. LOAD   - read every .md file in the knowledge/ folder
  2. SPLIT  - cut the text into small chunks
  3. STORE  - turn each chunk into a vector (a list of numbers) and keep them in FAISS, the vector database
  4. ANSWER - find the chunks most similar to the question and hand them, plus the chat so far, to the LLM
  5. LEADS  - appointment requests are saved in a Postgres database, emailed to you, and downloadable as an Excel sheet
"""
import os
import secrets
from io import BytesIO
from pathlib import Path
from typing import Literal

import httpx
import psycopg
from dotenv import load_dotenv
from fastapi import BackgroundTasks, Depends, FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, Response
from fastapi.security import HTTPBasic, HTTPBasicCredentials
from langchain_community.embeddings import FastEmbedEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_groq import ChatGroq
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import Workbook
from pydantic import BaseModel, Field

load_dotenv()  # reads the keys from the .env file when running on your computer

# 1. LOAD
docs = [Document(page_content=f.read_text(encoding="utf-8"), metadata={"source": f.name})
        for f in Path("knowledge").glob("*.md")]

# 2. SPLIT
chunks = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=100).split_documents(docs)

# 3. STORE (free embeddings that run locally, no API key needed)
# ponytail: index is rebuilt in memory on every start (a few seconds for a portfolio); use vector_db.save_local() if knowledge/ gets large
vector_db = FAISS.from_documents(chunks, FastEmbedEmbeddings(cache_dir="models"))
retriever = vector_db.as_retriever(search_kwargs={"k": 4})

# 4. ANSWER
prompt = ChatPromptTemplate.from_messages([
    ("system",
     "You are the friendly assistant on the portfolio website of Shohanur Rahman (Shohan), proshohan.com. "
     "Answer visitors' questions using ONLY the context below and the conversation so far. Refer to Shohan by name. "
     "Be professional, warm and brief (2-4 sentences). "
     "If the answer is not in the context, say you don't know and suggest emailing dev.shohanur@gmail.com. "
     "If the visitor wants to hire Shohan, get a quote, or book a call or meeting, invite them to fill in "
     "the appointment form and end your reply with [BOOK].\n\n"
     "Context:\n{context}"),
    MessagesPlaceholder("history"),  # the chat so far, so follow-up questions make sense
    ("human", "{question}"),
])
llm = ChatGroq(model=os.getenv("MODEL", "openai/gpt-oss-120b"), temperature=0.2)
chain = prompt | llm | StrOutputParser()


class Turn(BaseModel):
    role: Literal["user", "assistant"]
    content: str = Field(max_length=2000)


class Question(BaseModel):
    message: str = Field(min_length=1, max_length=500)
    history: list[Turn] = Field(default=[], max_length=10)


def ask(q: Question) -> dict:
    # search with the previous question too, so "tell me more about it" still finds the right chunks
    # ponytail: cheap follow-up trick; switch to an LLM query rewrite (history-aware retriever) if follow-ups still miss
    previous = next((t.content for t in reversed(q.history) if t.role == "user"), "")
    found = retriever.invoke(f"{previous}\n{q.message}")
    context = "\n\n".join(d.page_content for d in found)
    history = [(t.role, t.content) for t in q.history]
    reply = chain.invoke({"context": context, "history": history, "question": q.message})
    return {"reply": reply.replace("[BOOK]", "").strip(), "book": "[BOOK]" in reply}


# 5. LEADS
class Lead(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    email: str = Field(max_length=200, pattern=r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
    phone: str = Field("", max_length=40)
    service: str = Field(min_length=1, max_length=100)
    preferred_time: str = Field(min_length=1, max_length=80)
    details: str = Field("", max_length=1000)


def database():
    conn = psycopg.connect(os.environ["DATABASE_URL"])
    conn.execute("""CREATE TABLE IF NOT EXISTS leads (
        id SERIAL PRIMARY KEY, created_at TIMESTAMPTZ DEFAULT now(),
        name TEXT, email TEXT, phone TEXT, service TEXT, preferred_time TEXT, details TEXT)""")
    return conn


def notify(lead: Lead):
    """Email you about a new booking via resend.com (Render's free plan blocks normal SMTP email)."""
    key = os.getenv("RESEND_API_KEY")
    if not key:
        return  # email is optional: the booking is already saved in the database
    details = "\n".join(f"{field.replace('_', ' ').title()}: {value or '-'}" for field, value in lead.model_dump().items())
    httpx.post("https://api.resend.com/emails", timeout=10, headers={"Authorization": f"Bearer {key}"}, json={
        "from": "Portfolio Chatbot <onboarding@resend.dev>",
        "to": os.getenv("NOTIFY_EMAIL", "dev.shohanur@gmail.com"),  # must be the email you signed up to Resend with
        "reply_to": lead.email,  # press Reply in your inbox to answer the visitor directly
        "subject": f"📅 New appointment request from {lead.name}",
        "text": f"{details}\n\nAll requests: download leads.xlsx from your chatbot.",
    }).raise_for_status()


def to_excel(headers: list, rows: list) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":  # a visitor typed text starting with "=": keep it as text, never run it as a formula
                cell.data_type = "s"
    for column in "ABCDEFG":
        sheet.column_dimensions[column].width = 24
    out = BytesIO()
    book.save(out)
    return out.getvalue()


# The web API that the chat widget on your site talks to
app = FastAPI()
app.add_middleware(CORSMiddleware, allow_methods=["POST"], allow_headers=["*"],
                   allow_origins=["https://proshohan.com", "https://www.proshohan.com"])


@app.post("/chat")
def chat(q: Question):
    return ask(q)


@app.post("/lead")
def save_lead(lead: Lead, background: BackgroundTasks):
    with database() as conn:
        conn.execute("INSERT INTO leads (name, email, phone, service, preferred_time, details) VALUES "
                     "(%(name)s, %(email)s, %(phone)s, %(service)s, %(preferred_time)s, %(details)s)",
                     lead.model_dump())
    background.add_task(notify, lead)  # sent after the visitor sees "Thanks", so they never wait for email
    return {"ok": True}


@app.get("/leads.xlsx")
def download_leads(login: HTTPBasicCredentials = Depends(HTTPBasic())):  # the browser asks for a password
    password = os.getenv("ADMIN_PASSWORD")
    if not password or not secrets.compare_digest(login.password.encode(), password.encode()):
        raise HTTPException(401, "Wrong password", headers={"WWW-Authenticate": "Basic"})
    with database() as conn:
        cur = conn.execute("SELECT created_at AT TIME ZONE INTERVAL '+06:00' AS received_dhaka_time, name, email, phone, "
                           "service, preferred_time, details FROM leads ORDER BY id DESC")
        rows = cur.fetchall()
        headers = [c.name for c in cur.description]
    return Response(to_excel(headers, rows),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="leads.xlsx"'})


@app.get("/widget.js")
def widget():
    return FileResponse("widget.js", media_type="application/javascript")


@app.get("/")
def home():  # test page: open it in a browser to try the widget
    return HTMLResponse('<meta name="viewport" content="width=device-width,initial-scale=1">'
                        '<h1>Portfolio chatbot is running</h1><script src="/widget.js" defer></script>')
